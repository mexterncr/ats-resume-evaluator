from flask import Flask, render_template, request, jsonify, session
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import generate_password_hash, check_password_hash
from pypdf import PdfReader
import docx
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
import re
import os
import random
import requests
from datetime import datetime, timedelta
from google.oauth2 import id_token
from google.auth.transport import requests as google_requests

app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'skillsync-secret-key-prod-2026')

raw_db_url = os.environ.get('DATABASE_URL', 'sqlite:///skillsync.db')
if raw_db_url.startswith("postgres://"):
    raw_db_url = raw_db_url.replace("postgres://", "postgresql://", 1)

app.config['SQLALCHEMY_DATABASE_URI'] = raw_db_url
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

# Credentials
SMTP_EMAIL = os.environ.get("SMTP_EMAIL", "huzaifayhchannel@gmail.com")
BREVO_API_KEY = os.environ.get("BREVO_API_KEY", "")
GOOGLE_CLIENT_ID = os.environ.get("GOOGLE_CLIENT_ID", "1063711450384-jqhqdt2igs7eldt2ldsms9ug55skrj4g.apps.googleusercontent.com")

EMAIL_REGEX = r'^[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z]{2,}$'

db = SQLAlchemy(app)

class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    email = db.Column(db.String(120), unique=True, nullable=False)
    password_hash = db.Column(db.String(256), nullable=True)
    is_verified = db.Column(db.Boolean, default=False)
    otp_code = db.Column(db.String(6), nullable=True)
    otp_expiry = db.Column(db.DateTime, nullable=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    scans = db.relationship('ScanHistory', backref='owner', lazy=True)

class ScanHistory(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    filename = db.Column(db.String(256), nullable=False)
    candidate_name = db.Column(db.String(128), default="Candidate")
    ats_score = db.Column(db.Float, nullable=False)
    status = db.Column(db.String(64), nullable=False)
    similarity = db.Column(db.Float, nullable=False)
    skill_score = db.Column(db.Float, nullable=False)
    sec_score = db.Column(db.Float, nullable=False)
    role_title = db.Column(db.String(128), default="Custom Role")
    timestamp = db.Column(db.String(64), nullable=False)

with app.app_context():
    db.create_all()

TECH_SKILLS_DB = {
    "python", "java", "c++", "c", "javascript", "typescript", "html", "css", "sql", 
    "mysql", "postgresql", "mongodb", "react", "angular", "vue", "django", "flask", 
    "fastapi", "spring boot", "node.js", "express", "aws", "docker", "kubernetes", 
    "git", "github", "linux", "rest api", "machine learning", "deep learning", "nlp",
    "scikit-learn", "tensorflow", "pytorch", "pandas", "numpy", "data structures"
}

SECTIONS_TO_CHECK = {
    "Contact Info": [r"email", r"phone", r"linkedin", r"github"],
    "Education": [r"education", r"academic", r"bachelor", r"degree", r"university"],
    "Technical Skills": [r"skills", r"technical skills", r"technologies"],
    "Projects": [r"projects", r"academic projects", r"portfolio"],
    "Experience": [r"experience", r"employment", r"internship"],
    "Certifications": [r"certifications", r"courses", r"achievements"]
}

SUPPORTED_EXTENSIONS = ('.pdf', '.docx', '.doc', '.txt')

def clean_text(text):
    text = text.lower()
    text = re.sub(r'[^a-zA-Z0-9\s]', ' ', text)
    return re.sub(r'\s+', ' ', text).strip()

def extract_candidate_name(raw_text, filename):
    if not raw_text:
        return "Candidate"

    DISQUALIFIERS = {
        "curriculum", "vitae", "resume", "profile", "contact", "email", "phone",
        "developer", "engineer", "scientist", "analyst", "intern", "lead", "manager",
        "mumbai", "maharashtra", "delhi", "bangalore", "pune", "india", "address",
        "education", "skills", "projects", "experience", "certifications", "summary",
        "road", "nagar", "lane", "street", "pincode", "github", "linkedin"
    }

    # Line-by-line inspection of document header
    raw_lines = [line.strip() for line in raw_text.splitlines() if line.strip()]

    for line in raw_lines[:12]:
        # Skip contact numbers, links, addresses or emails
        if any(char in line for char in ["@", "http", "www.", "/", "+", ":", "|", "#"]):
            continue
        if re.search(r'\d', line):
            continue

        clean_tokens = [w for w in re.split(r'[^a-zA-Z]', line) if w]
        
        # Valid names are 2 to 4 alphabetic words
        if 2 <= len(clean_tokens) <= 4:
            lowered_tokens = [t.lower() for t in clean_tokens]
            if any(token in DISQUALIFIERS for token in lowered_tokens):
                continue
            return " ".join(clean_tokens).title()

    # Fallback to cleaned filename if document header has unconventional formatting
    clean_fn = re.sub(r'\.[^/.]+$', '', filename)
    clean_fn = re.sub(r'\(\d+\)', '', clean_fn)
    fn_first_part = re.split(r'[-_]', clean_fn)[0].strip()
    fn_tokens = [w for w in fn_first_part.split() if w.lower() not in DISQUALIFIERS and not re.search(r'\d', w)]
    if len(fn_tokens) >= 2:
        return " ".join(fn_tokens).title()

    return "Candidate"

def extract_file_content(file):
    filename = file.filename.lower()
    extracted_text = ""
    try:
        file.seek(0)
        if filename.endswith('.pdf'):
            reader = PdfReader(file)
            for page in reader.pages:
                txt = page.extract_text()
                if txt:
                    extracted_text += txt + "\n"
        elif filename.endswith(('.docx', '.doc')):
            doc = docx.Document(file)
            for para in doc.paragraphs:
                if para.text:
                    extracted_text += para.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text:
                            extracted_text += cell.text + " "
        elif filename.endswith('.txt'):
            extracted_text = file.read().decode('utf-8', errors='ignore')
    except Exception as e:
        print(f"Error reading {file.filename}: {e}")
        extracted_text = ""
    return extracted_text

def send_real_email_otp(recipient_email, otp_code, purpose):
    clean_key = (BREVO_API_KEY or "").strip()
    if not clean_key:
        return False, "BREVO_API_KEY not configured in server environment."

    url = "https://api.brevo.com/v3/smtp/email"
    headers = {
        "accept": "application/json",
        "api-key": clean_key,
        "content-type": "application/json"
    }
    payload = {
        "sender": {"name": "SkillSync AI", "email": SMTP_EMAIL},
        "to": [{"email": recipient_email}],
        "subject": f"SkillSync AI — Verification Code: {otp_code}",
        "htmlContent": f"""
            <div style="font-family: Arial, sans-serif; padding: 24px; color: #1e293b; max-width: 500px; margin: 0 auto; border: 1px solid #e2e8f0; border-radius: 12px;">
              <h2 style="color: #2563eb; margin-bottom: 8px;">SkillSync AI Verification</h2>
              <p style="font-size: 14px; color: #475569;">Use the following 6-digit verification code to complete your {purpose} request:</p>
              <div style="font-size: 32px; font-weight: bold; letter-spacing: 6px; color: #1e293b; background-color: #f1f5f9; padding: 14px; text-align: center; border-radius: 8px; margin: 20px 0;">
                {otp_code}
              </div>
              <p style="font-size: 12px; color: #94a3b8;">This code is valid for 10 minutes. If you did not request this, please ignore.</p>
            </div>
        """
    }

    try:
        response = requests.post(url, json=payload, headers=headers, timeout=10)
        return response.status_code in [200, 201, 202], "Dispatched"
    except Exception as e:
        return False, str(e)

@app.route('/')
def home():
    return render_template('index.html', google_client_id=GOOGLE_CLIENT_ID)

@app.route('/api/auth/status', methods=['GET'])
def auth_status():
    user_id = session.get('user_id')
    if user_id:
        user = User.query.get(user_id)
        if user:
            scan_count = ScanHistory.query.filter_by(user_id=user.id).count()
            return jsonify({'logged_in': True, 'email': user.email, 'scan_count': scan_count})
    return jsonify({'logged_in': False})

@app.route('/api/auth/send-otp', methods=['POST'])
def send_otp():
    data = request.get_json() or {}
    email = data.get('email', '').strip().lower()
    purpose = data.get('purpose', 'register')

    if not email or not re.match(EMAIL_REGEX, email):
        return jsonify({'error': 'Please enter a valid email address (e.g. name@gmail.com).'}), 400

    user = User.query.filter_by(email=email).first()
    if purpose == 'register' and user and user.is_verified:
        return jsonify({'error': 'An account with this email already exists.'}), 400
    if purpose == 'forgot' and not user:
        return jsonify({'error': 'No account found with this email address.'}), 404

    code = f"{random.randint(100000, 999999)}"
    expiry = datetime.utcnow() + timedelta(minutes=10)

    if not user:
        user = User(email=email, password_hash="pending", is_verified=False)
        db.session.add(user)

    user.otp_code = code
    user.otp_expiry = expiry
    db.session.commit()

    email_sent, err_msg = send_real_email_otp(email, code, purpose)

    if email_sent:
        return jsonify({'success': True, 'message': f'Verification code sent to {email}. Check your inbox!'})
    else:
        return jsonify({'error': f'Failed to dispatch email: {err_msg}'}), 400

@app.route('/api/auth/verify-register', methods=['POST'])
def verify_register():
    data = request.get_json() or {}
    email = data.get('email', '').strip().lower()
    otp = data.get('otp', '').strip()
    password = data.get('password', '').strip()

    if not email or not re.match(EMAIL_REGEX, email):
        return jsonify({'error': 'Please enter a valid email address.'}), 400
    if not otp or not password:
        return jsonify({'error': 'OTP and password are required.'}), 400
    if len(password) < 6:
        return jsonify({'error': 'Password must be at least 6 characters.'}), 400

    user = User.query.filter_by(email=email).first()
    if not user or user.otp_code != otp or datetime.utcnow() > (user.otp_expiry or datetime.min):
        return jsonify({'error': 'Invalid or expired verification code.'}), 400

    user.password_hash = generate_password_hash(password)
    user.is_verified = True
    user.otp_code = None
    user.otp_expiry = None
    db.session.commit()

    session['user_id'] = user.id
    return jsonify({'success': True, 'email': user.email, 'scan_count': 0})

@app.route('/api/auth/reset-password', methods=['POST'])
def reset_password():
    data = request.get_json() or {}
    email = data.get('email', '').strip().lower()
    otp = data.get('otp', '').strip()
    new_password = data.get('new_password', '').strip()

    if not email or not re.match(EMAIL_REGEX, email):
        return jsonify({'error': 'Please enter a valid email address.'}), 400
    if not otp or not new_password:
        return jsonify({'error': 'All fields are required.'}), 400
    if len(new_password) < 6:
        return jsonify({'error': 'Password must be at least 6 characters.'}), 400

    user = User.query.filter_by(email=email).first()
    if not user or user.otp_code != otp or datetime.utcnow() > (user.otp_expiry or datetime.min):
        return jsonify({'error': 'Invalid or expired verification code.'}), 400

    user.password_hash = generate_password_hash(new_password)
    user.otp_code = None
    user.otp_expiry = None
    db.session.commit()

    return jsonify({'success': True, 'message': 'Password reset successful. Please login now.'})

@app.route('/api/auth/google', methods=['POST'])
def google_auth():
    data = request.get_json() or {}
    token = data.get('credential')

    if not token:
        return jsonify({'error': 'Missing Google token.'}), 400

    try:
        id_info = id_token.verify_oauth2_token(token, google_requests.Request(), GOOGLE_CLIENT_ID)
        email = id_info['email'].lower()

        user = User.query.filter_by(email=email).first()
        if not user:
            user = User(email=email, password_hash=generate_password_hash(os.urandom(16).hex()), is_verified=True)
            db.session.add(user)
            db.session.commit()

        session['user_id'] = user.id
        scan_count = ScanHistory.query.filter_by(user_id=user.id).count()
        return jsonify({'success': True, 'email': user.email, 'scan_count': scan_count})
    except Exception as e:
        return jsonify({'error': f'Google authentication failed: {str(e)}'}), 400

@app.route('/api/auth/login', methods=['POST'])
def login():
    data = request.get_json() or {}
    email = data.get('email', '').strip().lower()
    password = data.get('password', '').strip()

    if not email or not re.match(EMAIL_REGEX, email):
        return jsonify({'error': 'Please enter a valid email address.'}), 400

    user = User.query.filter_by(email=email).first()
    if not user or not user.is_verified or not check_password_hash(user.password_hash, password):
        return jsonify({'error': 'Invalid email or password credentials.'}), 401

    session['user_id'] = user.id
    scan_count = ScanHistory.query.filter_by(user_id=user.id).count()
    return jsonify({'success': True, 'email': user.email, 'scan_count': scan_count})

@app.route('/api/auth/logout', methods=['POST'])
def logout():
    session.pop('user_id', None)
    return jsonify({'success': True})

@app.route('/api/history', methods=['GET'])
def get_history():
    user_id = session.get('user_id')
    if not user_id:
        return jsonify({'error': 'Please login to access audit history.'}), 401

    scans = ScanHistory.query.filter_by(user_id=user_id).order_by(ScanHistory.id.desc()).all()
    history_data = [{
        'id': s.id,
        'filename': s.filename,
        'candidate_name': getattr(s, 'candidate_name', 'Candidate'),
        'ats_score': s.ats_score,
        'status': s.status,
        'similarity': s.similarity,
        'skill_score': s.skill_score,
        'sec_score': s.sec_score,
        'role_title': s.role_title,
        'timestamp': s.timestamp
    } for s in scans]

    return jsonify({'history': history_data})

@app.route('/analyze', methods=['POST'])
def analyze():
    job_desc = request.form.get('job_desc', '').strip()
    role_name = request.form.get('role_name', 'Technical Position').strip()
    resume_files = request.files.getlist('resume_file')
    is_demo = request.form.get('is_demo') == 'true'

    if not job_desc:
        return jsonify({'error': 'Please provide a target Job Description to continue.'}), 400

    clean_jd = clean_text(job_desc)
    resumes_data = []

    if is_demo:
        demo_text = """
        Mohammed Masiha. Email: student@email.com. Phone: +91 9876543210.
        Education: Bachelor of Science in Computer Science, 2026.
        Technical Skills: Python, SQL, C++, HTML, CSS, JavaScript, Flask, Git, GitHub, Machine Learning, Data Structures.
        Projects: AI Resume ATS Evaluator using Python NLP, Web Inventory System using Node.js and SQL.
        Certifications: Python Data Science.
        """
        resumes_data.append(("Sample_CS_Resume.pdf", demo_text))
    else:
        valid_files = [f for f in resume_files if f and f.filename.lower().endswith(SUPPORTED_EXTENSIONS)]
        if not valid_files:
            return jsonify({'error': 'Please upload valid resume documents (.PDF, .DOCX, or .TXT).'}), 400

        for file in valid_files:
            text = extract_file_content(file)
            if not text.strip():
                text = f"Candidate Document: {file.filename}. Format parsed with minimal extractable tokens."
            resumes_data.append((file.filename, text))

    results = []
    vectorizer = TfidfVectorizer(stop_words='english', ngram_range=(1, 2))
    current_user_id = session.get('user_id')
    now_str = datetime.now().strftime('%d %b %Y, %I:%M %p')

    for fname, r_text in resumes_data:
        clean_res = clean_text(r_text)
        cand_name = extract_candidate_name(r_text, fname)

        try:
            tfidf_matrix = vectorizer.fit_transform([clean_jd, clean_res])
            similarity = round(float(cosine_similarity(tfidf_matrix[0:1], tfidf_matrix[1:2])[0][0]) * 100, 1)
        except Exception:
            similarity = 0.0

        sections = {}
        for name, kws in SECTIONS_TO_CHECK.items():
            sections[name] = any(re.search(rf"\b{kw}\b", clean_res) for kw in kws)
        sec_score = round((sum(sections.values()) / len(sections)) * 100, 1)

        jd_skills = {s for s in TECH_SKILLS_DB if re.search(rf"\b{re.escape(s)}\b", clean_jd)}
        res_skills = {s for s in TECH_SKILLS_DB if re.search(rf"\b{re.escape(s)}\b", clean_res)}
        
        matched = sorted(list(jd_skills.intersection(res_skills)))
        missing = sorted(list(jd_skills.difference(res_skills)))
        skill_score = round((len(matched) / len(jd_skills) * 100), 1) if jd_skills else 100.0

        ats_score = round((similarity * 0.40) + (skill_score * 0.40) + (sec_score * 0.20), 1)

        if ats_score >= 60.0:
            status = "Shortlisted"
        elif ats_score >= 45.0:
            status = "Potential Match"
        else:
            status = "Rejected"

        if current_user_id:
            scan_record = ScanHistory(
                user_id=current_user_id,
                filename=fname,
                candidate_name=cand_name,
                ats_score=ats_score,
                status=status,
                similarity=similarity,
                skill_score=skill_score,
                sec_score=sec_score,
                role_title=role_name or "Custom Role",
                timestamp=now_str
            )
            db.session.add(scan_record)

        results.append({
            'filename': fname,
            'candidate_name': cand_name,
            'timestamp': now_str,
            'ats_score': ats_score,
            'status': status,
            'similarity': similarity,
            'skill_score': skill_score,
            'sec_score': sec_score,
            'matched_skills': matched,
            'missing_skills': missing,
            'sections': sections
        })

    if current_user_id:
        db.session.commit()

    results = sorted(results, key=lambda x: x['ats_score'], reverse=True)

    if len(results) == 1:
        return jsonify(results[0])
    return jsonify({'multiple': True, 'results': results})

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)
